# -*- coding: utf-8 -*-
"""Чтение файлов разных форматов."""
import os

from lektor.textproc.files import read_file, SUPPORTED


def test_html_stripped(tmp_path):
    p = tmp_path / "страница.html"
    p.write_text(
        "<html><head><style>.x{}</style><script>var a=1;</script></head>"
        "<body><h1>Заголовок</h1><p>Первый <b>абзац</b> текста.</p>"
        "<ul><li>пункт раз</li><li>пункт два</li></ul></body></html>",
        encoding="utf-8")
    t = read_file(str(p))
    flat = " ".join(t.split())
    assert "var a" not in t and ".x" not in t
    assert "Заголовок" in flat
    assert "Первый абзац текста." in flat   # инлайн-теги не рвут фразу
    assert ".html" in SUPPORTED


def test_txt_encodings(tmp_path):
    p = tmp_path / "cp1251.txt"
    p.write_bytes("Проверка кодировки.".encode("cp1251"))
    assert "Проверка" in read_file(str(p))
    p2 = tmp_path / "utf16.txt"
    p2.write_bytes("Другая кодировка.".encode("utf-16"))
    assert "Другая" in read_file(str(p2))


def test_broken_docx_human_message(tmp_path):
    """Повреждённый .docx — понятное сообщение вместо техжаргона."""
    import pytest
    from lektor.textproc.files import read_file
    bad = tmp_path / "битый.docx"
    bad.write_text("это не docx", encoding="utf-8")
    with pytest.raises(RuntimeError) as e:
        read_file(str(bad))
    assert "повреждён" in str(e.value) or "docx" in str(e.value)


def test_pdf_scan_no_text_layer(tmp_path):
    """PDF без текстового слоя — подсказка про скан/OCR."""
    import pytest
    from lektor.textproc.files import read_file
    try:
        import pymupdf as fitz  # noqa
    except ImportError:
        pytest.skip("PyMuPDF не установлен")
    # PyMuPDF 1.28 кэширует sys.stdout при импорте и падает в закрытый
    # дескриптор под capsys — подменяем его буфер на вечный
    try:
        import sys
        if getattr(fitz, "utils", None) is not None and \
                hasattr(fitz, "_g_out_message"):
            fitz._g_out_message = sys.stderr
    except Exception:
        pass
    doc = fitz.Document()
    doc.new_page()
    scan = tmp_path / "скан.pdf"
    doc.save(str(scan))
    with pytest.raises(RuntimeError) as e:
        read_file(str(scan))
    assert "текстового слоя" in str(e.value) or "скан" in str(e.value)


def test_ssml_broken_tags_survive():
    """Битые/чужие SSML-теги не ломают чтение — удаляются, текст остаётся."""
    from lektor.textproc.ssml import extract_ssml
    from lektor.report import Report
    r = Report()
    res = extract_ssml('Текст <break time="abc"> мусор и <speak>незакрытый.', r)
    assert "мусор" in res.text and "незакрытый" in res.text
    assert "<" not in res.text
    res2 = extract_ssml('<script>alert(1)</script> после тега.', r)
    # сам тег исчезает; содержимое тега остаётся текстом (не ломает конвейер)
    assert "после тега" in res2.text


def test_read_rtf_strips_commands():
    """RTF читается парсером: команды вырезаются, \\'XX и \\uNNNN — буквы."""
    import os
    import tempfile
    from lektor.textproc.files import read_file
    rtf = (r"{\rtf1\ansi\ansicpg1251{\fonttbl{\f0 Times;}}"
        r"\par\b \u1055?\u1088?\u1080?\u1075?\u1086?\u1074?\u1086?\u1088;?\b0  "
           r"\'e8\'e2\'e0\'ed \'c8\'e2\'e0\'ed \par"
           r"\par \'d1\'f2\'e0\'f2\'fc\'ff 158 \\'d3\'ca.\par}")
    with tempfile.NamedTemporaryFile(suffix=".rtf", delete=False, mode="w",
                                     encoding="ascii") as tf:
        tf.write(rtf)
        tmp = tf.name
    try:
        out = read_file(tmp)
        assert "rtf1" not in out and "fonttbl" not in out and "\\" not in out
        assert "Приговор" in out and "иван" in out.lower()
        assert "Статья 158 УК." in out
    finally:
        os.unlink(tmp)


def test_docx_table_order():
    """Таблица в docx читается на своём месте, а не выпадает в конец."""
    import os
    import tempfile
    from lektor.textproc.files import read_file
    from docx import Document
    d = Document()
    d.add_paragraph("Первый абзац.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Ячейка А"
    t.rows[0].cells[1].text = "Ячейка Б"
    d.add_paragraph("Последний абзац.")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
        d.save(tf.name)
        tmp = tf.name
    try:
        out = read_file(tmp)
        first = out.find("Ячейка А")
        assert 0 <= out.find("Первый") < first < out.find("Последний")
    finally:
        os.unlink(tmp)


FB2_XML = ('<?xml version="1.0"?>\n'
           '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">'
           '<body><section><p>Незазипованный fb2.</p>'
           '<p>Вторая строка книги.</p></section></body></FictionBook>')


def test_fb2_plain_xml(tmp_path):
    """fb2 без zip-обёртки (обычный XML) читается как книга."""
    p = tmp_path / "книга.fb2"
    p.write_text(FB2_XML, encoding="utf-8")
    out = read_file(str(p))
    assert "Незазипованный fb2." in out
    assert "Вторая строка книги." in out


def test_fb2_zipped(tmp_path):
    """Зазипованный fb2 по-прежнему читается из zip-контейнера."""
    import zipfile
    p = tmp_path / "книга_z.fb2"
    with zipfile.ZipFile(str(p), "w") as z:
        z.writestr("book.fb2", FB2_XML)
    out = read_file(str(p))
    assert "Незазипованный fb2." in out


def test_epub_garbage_clear_error(tmp_path):
    """Мусор с расширением .epub — внятная ошибка, не сырой BadZipFile."""
    import pytest
    p = tmp_path / "битый.epub"
    p.write_bytes(b"not a zip at all")
    with pytest.raises(RuntimeError) as e:
        read_file(str(p))
    assert "epub" in str(e.value)
